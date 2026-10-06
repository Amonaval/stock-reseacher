import io
import re
import pandas as pd
from collections import defaultdict


def norm_company(v):
    s = re.sub(r"\s+", " ", str(v or "").strip())
    return s.casefold()


def _find_col(columns, names):
    lookup = {str(c).strip().casefold(): c for c in columns}
    for name in names:
        if name.casefold() in lookup:
            return lookup[name.casefold()]
    return None


def read_result_file(uploaded_file, strategy_id=None):
    raw = uploaded_file.getvalue()
    name = uploaded_file.name.lower()
    hyperlink_by_row = {}
    if name.endswith(".xlsx"):
        df = pd.read_excel(io.BytesIO(raw))
        # pandas drops cell hyperlinks. Preserve exact Screener company URLs in
        # the fallback/debug importer rather than searching again by company name.
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw), read_only=False, data_only=True)
            ws = wb.active
            headers = [str(c.value or "").strip() for c in ws[1]]
            company_idx = next((i for i,h in enumerate(headers,1) if h.casefold() in {"name","company","company name","stock","symbol"}), None)
            if company_idx:
                for excel_row in range(2, ws.max_row + 1):
                    cell = ws.cell(excel_row, company_idx)
                    if cell.hyperlink and cell.hyperlink.target:
                        hyperlink_by_row[excel_row - 2] = str(cell.hyperlink.target)
        except Exception:
            hyperlink_by_row = {}
    elif name.endswith(".xls"):
        df = pd.read_excel(io.BytesIO(raw))
    elif name.endswith(".csv"):
        df = pd.read_csv(io.BytesIO(raw))
    else:
        raise ValueError("Upload a CSV/XLS/XLSX Screener result export.")

    company_col = _find_col(df.columns, ["name", "company", "company name", "stock", "symbol"])
    if company_col is None:
        for c in df.columns:
            if df[c].dtype == object:
                company_col = c
                break
    if company_col is None:
        raise ValueError("Could not identify the company/name column in the result file.")

    sid_col = _find_col(df.columns, ["strategy_id", "strategy", "screen_id"])
    url_col = _find_col(df.columns, ["url", "company url", "link"])

    out = []
    for row_index, (_, row) in enumerate(df.iterrows()):
        company = str(row.get(company_col, "") or "").strip()
        if not company or company.lower() == "nan":
            continue
        sid = str(row.get(sid_col, "") or "").strip() if sid_col else str(strategy_id or "").strip()
        snapshot = {}
        for col in df.columns:
            value = row.get(col)
            if pd.isna(value):
                value = None
            elif hasattr(value, "item"):
                try:
                    value = value.item()
                except Exception:
                    pass
            snapshot[str(col)] = value
        out.append({
            "company": company,
            "company_key": norm_company(company),
            "strategy_id": sid,
            "url": (str(row.get(url_col, "") or "").strip() if url_col else "") or hyperlink_by_row.get(row_index, ""),
            "source_file": uploaded_file.name,
            "snapshot": snapshot,
        })
    return out


def parse_pasted_companies(text, strategy_id):
    rows = []
    for line in str(text or "").splitlines():
        company = line.strip().strip(",")
        if not company:
            continue
        rows.append({"company": company,"company_key": norm_company(company),"strategy_id": strategy_id,"url": "","source_file": "manual_paste"})
    return rows


def build_candidate_universe(rows, strategies):
    strategy_map = {s["id"]: s for s in strategies}
    grouped = defaultdict(lambda: {"company": "", "urls": set(), "strategies": set(), "sources": set(), "snapshots": []})
    for row in rows:
        key = row.get("company_key") or norm_company(row.get("company"))
        if not key: continue
        g = grouped[key]
        if not g["company"]: g["company"] = row.get("company", "")
        if row.get("url"): g["urls"].add(row["url"])
        if row.get("strategy_id"): g["strategies"].add(row["strategy_id"])
        if row.get("source_file"): g["sources"].add(row["source_file"])
        if row.get("snapshot"): g["snapshots"].append(row["snapshot"])

    result = []
    max_possible = max(1, len(strategy_map))
    for key, g in grouped.items():
        ids = sorted(g["strategies"])
        confidence_sum = sum(strategy_map.get(sid, {}).get("evidence_confidence", 0) for sid in ids)
        average_conf = confidence_sum / max(1, len(ids))
        overlap = len(ids)
        priority = 100 * (0.70 * overlap / max_possible + 0.30 * average_conf / 100)
        result.append({
            "company": g["company"],"strategy_count": overlap,"strategies": ", ".join(ids),
            "strategy_names": ", ".join(strategy_map.get(x, {}).get("name", x) for x in ids),
            "methodology_confidence_avg": round(average_conf, 1),"research_priority_score": round(priority, 1),
            "url": sorted(g["urls"])[0] if g["urls"] else "","sources": ", ".join(sorted(g["sources"])),
            "snapshot": g["snapshots"][0] if g["snapshots"] else {},
        })
    result.sort(key=lambda x: (-x["strategy_count"], -x["research_priority_score"], x["company"].casefold()))
    return result
